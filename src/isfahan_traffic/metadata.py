from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, MutableMapping, Optional, Sequence, Tuple

import cv2
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .parser import discover_raw_month_files, iter_scats_records
from .paths import PROJECT_ROOT, ensure_project_directories, load_settings


APPROACH_ORDER = ("south", "west", "north", "east")


def _read_image_unicode(path: Path) -> np.ndarray:
    #read bytes first to support Persian file paths on Windows
    image = cv2.imdecode(np.fromfile(str(path), dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError("Could not decode diagram: {0}".format(path))
    return image


def detect_detector_boxes(path: Path) -> List[Tuple[float, float]]:

    image = _read_image_unicode(path)
    blue = (image[:, :, 0] == 255) & (image[:, :, 1] == 0) & (image[:, :, 2] == 0)
    green = (image[:, :, 0] == 0) & (image[:, :, 1] == 255) & (image[:, :, 2] == 0)
    mask = ((blue | green).astype(np.uint8) * 255)
    contours, _ = cv2.findContours(mask, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    candidates: List[Tuple[int, int, int, int]] = []
    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        area = cv2.contourArea(contour)
        #the first 370 pixels contain the legend and title
        if x > 370 and 14 <= width <= 35 and 14 <= height <= 35 and area > 450:
            candidates.append((x, y, width, height))

    centers: List[Tuple[float, float]] = []
    for x, y, width, height in sorted(candidates):
        center = (x + width / 2.0, y + height / 2.0)
        if not any(abs(center[0] - seen[0]) < 3 and abs(center[1] - seen[1]) < 3 for seen in centers):
            centers.append(center)
    return centers


def classify_approaches(
    centers: Sequence[Tuple[float, float]], center_x: float = 690, center_y: float = 347
) -> Counter:
    groups: Counter = Counter()
    for x, y in centers:
        dx = x - center_x
        dy = y - center_y
        if abs(dx) > abs(dy):
            groups["east" if dx > 0 else "west"] += 1
        else:
            groups["south" if dy > 0 else "north"] += 1
    return groups


def extract_registry(source_root: Path) -> pd.DataFrame:
    workbook_path = next(
        path
        for path in source_root.rglob("*.xlsx")
        if "scats" in path.name.lower() and not path.name.startswith("~$")
    )
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = max(workbook.worksheets, key=lambda sheet: sheet.max_row)
    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        try:
            intersection_id = int(row[0])
        except (TypeError, ValueError):
            continue
        rows.append(
            {
                "intersection_id": intersection_id,
                "name_fa": str(row[1]).strip() if row[1] else "",
                "gis_id": int(row[2]) if row[2] is not None else "",
                "registry_source": workbook_path.name,
            }
        )
    return pd.DataFrame(rows).drop_duplicates("intersection_id")


def profile_channels(raw_files: Sequence[Path]) -> pd.DataFrame:
    counters: MutableMapping[Tuple[int, int], Counter] = defaultdict(Counter)
    record_counts: Counter = Counter()
    for file_index, path in enumerate(raw_files, start=1):
        print("[metadata] profiling {0}/{1}: {2}".format(file_index, len(raw_files), path.name), flush=True)
        for record in iter_scats_records(path):
            record_counts[record.intersection_id] += 1
            for channel_id, value in record.channels.items():
                counter = counters[(record.intersection_id, channel_id)]
                counter["observations"] += 1
                if value is None:
                    counter["invalid"] += 1
                else:
                    counter["finite"] += 1
                    if value == 0:
                        counter["zero"] += 1
                    elif value > 0:
                        counter["positive"] += 1
                        counter["volume"] += value

    rows = []
    for (intersection_id, channel_id), counter in sorted(counters.items()):
        observations = counter["observations"]
        finite = counter["finite"]
        rows.append(
            {
                "intersection_id": intersection_id,
                "channel_id": channel_id,
                "intersection_records": record_counts[intersection_id],
                "observations": observations,
                "finite": finite,
                "invalid": counter["invalid"],
                "zero": counter["zero"],
                "positive": counter["positive"],
                "volume": counter["volume"],
                "finite_rate": finite / observations if observations else 0.0,
                "positive_rate": counter["positive"] / observations if observations else 0.0,
            }
        )
    return pd.DataFrame(rows)


def diagram_inventory(source_root: Path) -> pd.DataFrame:
    rows = []
    for path in sorted(source_root.rglob("*.png")):
        if not path.stem.isdigit():
            continue
        centers = detect_detector_boxes(path)
        groups = classify_approaches(centers)
        rows.append(
            {
                "intersection_id": int(path.stem),
                "diagram_file": path.name,
                "diagram_relative_path": str(path.relative_to(source_root)),
                "detector_boxes": len(centers),
                "south_boxes": groups["south"],
                "west_boxes": groups["west"],
                "north_boxes": groups["north"],
                "east_boxes": groups["east"],
            }
        )
    return pd.DataFrame(rows)


def build_detector_registry(channel_profile: pd.DataFrame, diagrams: pd.DataFrame) -> pd.DataFrame:
    diagram_by_id = diagrams.set_index("intersection_id").to_dict("index")
    rows = []
    for intersection_id, group in channel_profile.groupby("intersection_id"):
        group = group.copy()
        diagram = diagram_by_id.get(int(intersection_id))
        candidates = group.to_dict("records")
        if diagram and int(diagram["detector_boxes"]) > 0:
            configured_count = min(int(diagram["detector_boxes"]), len(candidates))
            #use the diagram count to select the most active channels
            #the small detector labels are too unreliable to read automatically
            ranked = sorted(
                candidates,
                key=lambda row: (row["positive"], row["finite"], -row["channel_id"]),
                reverse=True,
            )[:configured_count]
            configured_ids = sorted(int(row["channel_id"]) for row in ranked)
            slots: List[str] = []
            for approach in APPROACH_ORDER:
                slots.extend([approach] * int(diagram[approach + "_boxes"]))
            if len(slots) != len(configured_ids):
                slots = ["unknown"] * len(configured_ids)
            approach_by_channel = dict(zip(configured_ids, slots))
            verification = "diagram_auto"
        else:
            configured_ids = sorted(
                int(row["channel_id"])
                for row in candidates
                if int(row["positive"]) >= 100
            )
            approach_by_channel = {channel_id: "unknown" for channel_id in configured_ids}
            verification = "data_only"

        by_channel = {int(row["channel_id"]): row for row in candidates}
        for channel_id in configured_ids:
            stats = by_channel[channel_id]
            operational = int(stats["positive"]) >= 100
            rows.append(
                {
                    **stats,
                    "approach": approach_by_channel.get(channel_id, "unknown"),
                    "configured": True,
                    "operational": operational,
                    "verification": verification,
                }
            )

    detector_registry = pd.DataFrame(rows)
    override_path = PROJECT_ROOT / "config" / "channel_overrides.csv"
    overrides = pd.read_csv(override_path)
    if not detector_registry.empty:
        override_index = {
            (int(row.intersection_id), int(row.channel_id)): row
            for row in overrides.itertuples(index=False)
        }
        for index, row in detector_registry.iterrows():
            override = override_index.get((int(row["intersection_id"]), int(row["channel_id"])))
            if override is not None:
                detector_registry.at[index, "approach"] = override.approach
                detector_registry.at[index, "verification"] = override.verification
                detector_registry.at[index, "source_note"] = override.source_note
    return detector_registry.sort_values(["intersection_id", "channel_id"])


def build_metadata() -> None:
    ensure_project_directories()
    settings = load_settings()
    source_root: Path = settings["source_root"]
    metadata_dir = PROJECT_ROOT / "data" / "metadata"

    raw_files = discover_raw_month_files(source_root)
    registry = extract_registry(source_root)
    diagrams = diagram_inventory(source_root)
    channel_profile = profile_channels(raw_files)
    detectors = build_detector_registry(channel_profile, diagrams)

    observed_ids = pd.DataFrame(
        {"intersection_id": sorted(channel_profile["intersection_id"].unique())}
    )
    intersections = observed_ids.merge(registry, how="outer", on="intersection_id")
    intersections = intersections.merge(diagrams, how="outer", on="intersection_id")
    intersections["has_raw_data"] = intersections["intersection_id"].isin(observed_ids["intersection_id"])
    intersections["has_diagram"] = intersections["diagram_file"].notna()
    intersections["has_registry_name"] = intersections["name_fa"].fillna("").ne("")

    intersections.sort_values("intersection_id").to_csv(
        metadata_dir / "intersections.csv", index=False, encoding="utf-8-sig"
    )
    diagrams.sort_values("intersection_id").to_csv(
        metadata_dir / "diagram_summary.csv", index=False, encoding="utf-8-sig"
    )
    channel_profile.to_csv(
        metadata_dir / "channel_quality.csv", index=False, encoding="utf-8-sig"
    )
    detectors.to_csv(
        metadata_dir / "detector_channels.csv", index=False, encoding="utf-8-sig"
    )
    pd.DataFrame(
        {
            "source_file": [str(path.relative_to(source_root)) for path in raw_files],
            "source_month": [path.name for path in raw_files],
        }
    ).to_csv(metadata_dir / "source_files.csv", index=False, encoding="utf-8-sig")

    print(
        "[metadata] wrote {0} intersections, {1} diagram rows, and {2} configured channels".format(
            len(intersections), len(diagrams), len(detectors)
        ),
        flush=True,
    )
